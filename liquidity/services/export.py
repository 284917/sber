"""Выгрузка для администратора в Excel: аналитика, архив сделок, все заявки и предложения — одним файлом."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from liquidity import config
from liquidity.services.costs import offer_cost

HEADER_FILL = PatternFill("solid", fgColor="F2F4F5")
HEADER_FONT = Font(bold=True, color="3A3F45")
TITLE_FONT = Font(bold=True, size=14)
MONEY = "#,##0.00"
PERCENT = '0.00"%"'


def _sheet(wb, title, headers, rows, formats=None, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for col, fmt in (formats or {}).items():
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = fmt
    for i, header in enumerate(headers, 1):
        width = (widths or {}).get(i) or min(48, max(12, len(str(header)) + 2,
                                                      *(len(str(r[i - 1] or "")) + 2 for r in rows[:200])))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _date(value):
    return value.replace(microsecond=0) if value else None


def build_workbook(data, offers) -> bytes:
    """data — результат analytics.overview(db); offers — все предложения."""
    a = data["analytics"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"

    def head(*values):
        ws.append(list(values))
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL

    ws.append(["Платформа ликвидности · аналитика"])
    ws["A1"].font = TITLE_FONT
    ws.append([f"Сформировано {datetime.utcnow():%d.%m.%Y %H:%M} UTC"])
    ws.append([])
    for label, value in (("Заявок всего", a["total_apps"]), ("Предложений всего", a["total_offers"]),
                         ("Доля заявок, дошедших до сделки, %", a["deal_share"]),
                         ("Согласованных сделок", len(data["archive"]))):
        ws.append([label, value])
    ws.append([])
    head("Статус заявок", "Количество")
    for status, count in a["status_counts"].most_common():
        ws.append([config.APPLICATION_STATUSES.get(status, status), count])
    ws.append([])
    head("Валюта", "Заявок", "Средний размер", "Сумма")
    for row in a["avg_by_currency"]:
        ws.append([row["currency"], row["count"], round(row["avg"], 2), round(row["total"], 2)])
        for col in "CD":
            ws[f"{col}{ws.max_row}"].number_format = MONEY
    ws.append([])
    head("Воронка", "Заявок", "Доля, %")
    for stage in data["funnel"]:
        ws.append([stage["label"], stage["count"], stage["share"]])
    ws.append([])
    head("Спрос: страна получения", "Заявок")
    for code, count in a["top_demand"]:
        ws.append([config.COUNTRY_NAMES.get(code, code), count])
    ws.append([])
    head("Предложение: страна кредитора", "Предложений")
    for code, count in a["top_supply"]:
        ws.append([config.COUNTRY_NAMES.get(code, code), count])
    ws.column_dimensions["A"].width = 40
    for col in "BCD":
        ws.column_dimensions[col].width = 18

    deals = []
    for row in data["archive"]:
        app, offer = row["app"], row["offer"]
        cost = offer_cost(offer, app) if offer else None
        deals.append([app.id, _date(app.confirmed_at), app.borrower.display_name if app.borrower else None,
                      app.borrower.registration_number if app.borrower else None,
                      offer.sender.display_name if offer and offer.sender else None,
                      config.COUNTRY_NAMES.get(offer.sender.country_code) if offer and offer.sender else None,
                      app.currency_code, float(offer.funding_amount) if offer else float(app.funding_amount),
                      float(offer.annual_interest_rate) if offer else None,
                      float(offer.total_commission) if offer else None,
                      round(cost["effective_rate"], 2) if cost and cost["effective_rate"] is not None else None,
                      round(cost["overpayment"], 2) if cost and cost["overpayment"] is not None else None,
                      app.term_months, config.COUNTRY_NAMES.get(app.target_jurisdiction),
                      config.TRANSFER_METHODS.get(app.transfer_method), app.purpose,
                      row["confirmed_by"].email if row["confirmed_by"] else None])
    _sheet(wb, "Сделки", ["№", "Подтверждена", "Заёмщик", "ИНН", "Кредитор", "Страна кредитора", "Валюта", "Сумма",
                          "Ставка, %", "Комиссия, %", "Полная стоимость, %", "Переплата", "Срок, мес.",
                          "Страна получения", "Способ передачи", "Назначение", "Подтвердил"],
           deals, {2: "dd.mm.yyyy hh:mm", 8: MONEY, 9: PERCENT, 10: PERCENT, 11: PERCENT, 12: MONEY})

    apps = [[app.id, _date(app.created_at), config.APPLICATION_STATUSES.get(app.status, app.status),
             app.borrower.display_name if app.borrower else None, app.currency_code, float(app.funding_amount),
             app.term_months, config.COUNTRY_NAMES.get(app.target_jurisdiction),
             config.TRANSFER_METHODS.get(app.transfer_method), app.purpose, app.moderation_comment]
            for app in data["apps"]]
    _sheet(wb, "Заявки", ["№", "Создана", "Статус", "Заёмщик", "Валюта", "Сумма", "Срок, мес.", "Страна получения",
                          "Способ передачи", "Назначение", "Комментарий модератора"],
           apps, {2: "dd.mm.yyyy hh:mm", 6: MONEY})

    offer_rows = []
    for o in sorted(offers, key=lambda o: o.id):
        app = o.application
        cost = offer_cost(o, app) if app else None
        offer_rows.append([o.id, app.id if app else None, _date(o.created_at),
                           config.OFFER_STATUSES.get(o.status, o.status), o.sender.display_name if o.sender else None,
                           app.currency_code if app else None, float(o.funding_amount), float(o.annual_interest_rate),
                           float(o.lender_commission), float(o.total_commission),
                           round(cost["effective_rate"], 2) if cost and cost["effective_rate"] is not None else None,
                           o.hide_reason])
    _sheet(wb, "Предложения", ["№", "Заявка", "Отправлено", "Статус", "Кредитор", "Валюта", "Сумма", "Ставка, %",
                               "Комиссия кредитора, %", "Комиссия всего, %", "Полная стоимость, %", "Причина скрытия"],
           offer_rows, {3: "dd.mm.yyyy hh:mm", 7: MONEY, 8: PERCENT, 9: PERCENT, 10: PERCENT, 11: PERCENT})

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
