"""Калькулятор, профиль кредитора, сроки ответа, проект договора, выгрузка в Excel, сохранённые фильтры,
страница «Безопасность и документы»."""
import json
import re
from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

from conftest import PDF, Browser, login
from liquidity import models
from liquidity.database import SessionLocal
from liquidity.services import sla
from liquidity.services.creditors import rating
from test_features import make_company

FORM = {"currency": "EUR", "country": "DE", "term_months": "12", "transfer_method": "swift"}


def app_id_by(purpose):
    with SessionLocal() as db:
        return db.query(models.Application).filter_by(purpose=purpose).one().id


def test_landing_calculator_and_security_page():
    page = Browser().get("/")
    assert page.status_code == 200 and 'id="calc"' in page.text and 'id="r-effective"' in page.text
    data = json.loads(re.search(r"const data = (\{.*?\});", page.text).group(1))
    assert data["platform_commission"] == 3.0 and data["rates"]["EUR"] == 90 and data["rates_date"] == "27.10.2026"

    security = Browser().get("/security")
    assert security.status_code == 200
    for text in ("Кто что видит", "Выписка из", "Сертификат налогового резидентства", "24 ч"):
        assert text in security.text, text
    assert 'href="/security"' in page.text


def test_sla_status():
    now = datetime(2026, 10, 28, 12)
    assert sla.status(None, "offer") is None
    fresh = sla.status(now - timedelta(hours=3), "offer", now)
    assert not fresh["overdue"] and not fresh["soon"] and "осталось 21 ч" in fresh["label"]
    assert sla.status(now - timedelta(hours=21), "offer", now)["soon"]
    late = sla.status(now - timedelta(hours=30), "offer", now)
    assert late["overdue"] and late["label"] == "Просрочено на 6 ч"
    assert sla.status(now - timedelta(days=4), "kyc", now)["label"] == "Просрочено на 2 дн."
    assert sla.status(now - timedelta(hours=24, minutes=10), "offer", now)["label"] == "Срок ответа истёк"
    assert sla.overdue_count([now - timedelta(hours=30), now, None], "moderation", now) == 1


def test_rating_formula():
    assert rating({"acceptance": None, "deals_done": 0, "decided": 0, "response_hours": None}) is None
    assert rating({"acceptance": 0, "deals_done": 0, "decided": 1, "response_hours": 3}) is None  # мало данных
    top = rating({"acceptance": 100, "deals_done": 4, "decided": 4, "response_hours": 5})
    assert top["score"] == 5.0
    slow = rating({"acceptance": 0, "deals_done": 0, "decided": 3, "response_hours": 200})
    assert slow["score"] == 1.0


def test_creditor_profile_contract_draft_and_export():
    make_company("pr-borrower@test.ru", "borrower", "ООО «Профиль»", "RU", "7710140679")
    make_company("pr-other@test.ru", "borrower", "ООО «Чужой»", "RU", "7728168971")
    make_company("pr-creditor@test.cy", "creditor", "Profile Capital", "CY", "HE 11")
    make_company("pr-creditor2@test.ae", "creditor", "Other Capital", "AE", "AE 12")
    borrower, other, creditor, creditor2 = (login(e) for e in ("pr-borrower@test.ru", "pr-other@test.ru",
                                                                "pr-creditor@test.cy", "pr-creditor2@test.ae"))
    admin = login("admin@test.local", "admin-password")
    with SessionLocal() as db:
        creditor_id = db.query(models.User).filter_by(email="pr-creditor@test.cy").one().id

    borrower.post("/create_application", data={**FORM, "amount": "200000", "purpose": "Профиль кредитора"})
    app_id = app_id_by("Профиль кредитора")
    admin.post(f"/admin/application/{app_id}/moderate", data={"decision": "approve"})
    creditor.post(f"/create_offer/{app_id}", data={"amount": "200000", "rate": "8", "lender_comm": "1"})
    with SessionLocal() as db:
        offer_id = db.query(models.Offer).filter_by(application_id=app_id).one().id

    # Пока оператор не показал предложение, заёмщик профиль не видит
    assert borrower.get(f"/creditors/{creditor_id}").status_code == 403
    admin.post(f"/operator/offer/{offer_id}/show")
    page = borrower.get(f"/applications/{app_id}/offers").text
    assert f'href="/creditors/{creditor_id}"' in page and "Новый кредитор" in page
    profile = borrower.get(f"/creditors/{creditor_id}")
    assert profile.status_code == 200 and "Profile Capital" in profile.text and "Скорость ответа" in profile.text
    assert other.get(f"/creditors/{creditor_id}").status_code == 403
    assert creditor2.get(f"/creditors/{creditor_id}").status_code == 403
    assert creditor.get(f"/creditors/{creditor_id}").status_code == 200
    assert admin.get(f"/creditors/{creditor_id}").status_code == 200

    # Проект договора — только после принятия предложения и только участникам
    assert borrower.get(f"/deal/{app_id}/contract-draft.pdf").status_code in (403, 404)
    borrower.post(f"/accept_offer/{offer_id}")
    draft = borrower.get(f"/deal/{app_id}/contract-draft.pdf")
    assert draft.status_code == 200 and draft.content.startswith(b"%PDF") and len(draft.content) > 5000
    assert creditor.get(f"/deal/{app_id}/contract-draft.pdf").status_code == 200
    assert creditor2.get(f"/deal/{app_id}/contract-draft.pdf").status_code == 403
    assert "Проект договора" in borrower.get(f"/deal/{app_id}").text

    # Сделка согласована — рейтинг появляется
    creditor.post(f"/deal/{app_id}/upload_contract", files={"file": PDF})
    admin.post(f"/operator/deal/{app_id}/confirm")
    profile = borrower.get(f"/creditors/{creditor_id}").text
    assert "Рейтинг платформы" in profile and "из 5" in profile

    # Excel для администратора
    xlsx = admin.get("/admin/export.xlsx")
    assert xlsx.status_code == 200 and "attachment" in xlsx.headers["content-disposition"]
    wb = load_workbook(BytesIO(xlsx.content))
    assert wb.sheetnames == ["Аналитика", "Сделки", "Заявки", "Предложения"]
    deal_ids = [row[0].value for row in wb["Сделки"].iter_rows(min_row=2)]
    assert app_id in deal_ids
    assert borrower.get("/admin/export.xlsx").status_code == 403
    operator_like = login("pr-creditor@test.cy")
    assert operator_like.get("/admin/export.xlsx").status_code == 403


def test_saved_searches_notify_matching_creditors():
    make_company("ss-borrower@test.ru", "borrower", "ООО «Фильтр»", "RU", "7714617793")
    make_company("ss-creditor@test.cy", "creditor", "Filter Capital", "CY", "HE 21")
    make_company("ss-creditor2@test.ae", "creditor", "Filter Two", "AE", "AE 22")
    borrower, creditor, creditor2 = login("ss-borrower@test.ru"), login("ss-creditor@test.cy"), login("ss-creditor2@test.ae")
    admin = login("admin@test.local", "admin-password")

    assert creditor.post("/saved-searches", data={}).headers["location"].endswith("search_error=empty#saved")
    r = creditor.post("/saved-searches", data={"currency": "EUR", "amount_min": "1000000"})
    assert r.headers["location"].endswith("search_saved=1#saved")
    assert "search_error=duplicate" in creditor.post("/saved-searches", data={"currency": "EUR", "amount_min": "1000000"}).headers["location"]
    assert "search_error=range" in creditor.post("/saved-searches", data={"amount_min": "5", "amount_max": "1"}).headers["location"]
    creditor2.post("/saved-searches", data={"currency": "USD"})
    page = creditor.get("/dashboard?currency=EUR&amount_min=1000000").text
    assert "EUR · от 1" in page and "Уведомления о новых заявках" in page and "Этот фильтр уже сохранён" in page

    borrower.post("/create_application", data={**FORM, "amount": "2500000", "purpose": "Подходит под фильтр"})
    borrower.post("/create_application", data={**FORM, "amount": "500000", "purpose": "Слишком мала"})
    for purpose in ("Подходит под фильтр", "Слишком мала"):
        admin.post(f"/admin/application/{app_id_by(purpose)}/moderate", data={"decision": "approve"})
    match_id = app_id_by("Подходит под фильтр")
    with SessionLocal() as db:
        ids = {e: db.query(models.User).filter_by(email=e).one().id for e in ("ss-creditor@test.cy", "ss-creditor2@test.ae")}
        texts = [n.text for n in db.query(models.Notification).filter_by(user_id=ids["ss-creditor@test.cy"])]
        assert sum("по вашему фильтру" in t for t in texts) == 1 and any(f"№{match_id}" in t for t in texts)
        assert not db.query(models.Notification).filter(models.Notification.user_id == ids["ss-creditor2@test.ae"],
                                                        models.Notification.text.like("%по вашему фильтру%")).count()
        search_id = db.query(models.SavedSearch).filter_by(user_id=ids["ss-creditor@test.cy"]).one().id

    assert creditor2.post(f"/saved-searches/{search_id}/delete").status_code == 404  # чужой фильтр
    assert borrower.post("/saved-searches", data={"currency": "EUR"}).status_code == 403
    creditor.post(f"/saved-searches/{search_id}/delete")
    with SessionLocal() as db:
        assert db.get(models.SavedSearch, search_id) is None
